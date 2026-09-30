"""정상 수집 결과를 실행별로 보관하고 다음 실행의 비교 기준을 갱신."""

import argparse
import csv
import hashlib
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


FIELDS = ("run_stamp", "run_id", "file", "plans", "rs_violations", "rm_violations",
          "added", "removed", "changed", "sha256", "run_url")


def counts(report):
    wb = load_workbook(report, read_only=True, data_only=True)
    try:
        guide = wb["가이드위반"]
        rs = rm = 0
        group = None
        for cells in guide.iter_rows(min_col=2, max_col=3, values_only=True):
            number, carrier = cells
            if isinstance(number, str) and "가이드 위반" in number:
                group = "RS" if " RS " in number else "RM"
            elif isinstance(number, int) and carrier:
                if group == "RS":
                    rs += 1
                elif group == "RM":
                    rm += 1
        kinds = Counter()
        for (kind,) in wb["전일 대비 변동"].iter_rows(min_row=2, max_col=1, values_only=True):
            if kind:
                kinds["changed" if "변경" in kind else kind] += 1
        return {
            "plans": wb["홈페이지"].max_row - 1,
            "rs_violations": rs, "rm_violations": rm,
            "added": kinds["추가"], "removed": kinds["삭제"], "changed": kinds["changed"],
        }
    finally:
        wb.close()


def record(report, stamp, run_id, run_url, history):
    datetime.strptime(stamp, "%Y%m%d_%H%M%S")
    if not report.is_file():
        raise FileNotFoundError(report)
    stats = counts(report)
    if stats["plans"] == 0:
        raise RuntimeError("빈 보고서는 이력에 추가하지 않습니다")
    folder = history / stamp[:4] / stamp[4:6]
    folder.mkdir(parents=True, exist_ok=True)
    name = f"mvno_combined_plans_{stamp}_{run_id}.xlsx"
    target = folder / name
    if target.exists():
        raise FileExistsError(target)
    shutil.copy2(report, target)
    digest = hashlib.sha256(target.read_bytes()).hexdigest()
    index = history / "runs.csv"
    entry = {"run_stamp": stamp, "run_id": run_id, "file": target.as_posix(),
             **stats, "sha256": digest, "run_url": run_url}
    with index.open("a", encoding="utf-8-sig" if not index.exists() else "utf-8",
                    newline="") as file:
        writer = csv.DictWriter(file, fieldnames=FIELDS)
        if file.tell() == 0:
            writer.writeheader()
        writer.writerow(entry)
    shutil.copy2(target, history / "latest.xlsx")
    print(f"이력 저장: {target} / 요금제 {stats['plans']}건 / RS {stats['rs_violations']}건 / RM {stats['rm_violations']}건")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("report", type=Path)
    parser.add_argument("--stamp", required=True)
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--run-url", required=True)
    parser.add_argument("--history", type=Path, default=Path("history"))
    args = parser.parse_args()
    record(args.report, args.stamp, args.run_id, args.run_url, args.history)
