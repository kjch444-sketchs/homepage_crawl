"""가이드위반 시트를 메일 본문으로 보내고 위반 요금제를 텔레그램으로 알림."""

import argparse
import html
import json
import os
import smtplib
import urllib.parse
import urllib.request
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook


def guide_sections(report):
    wb = load_workbook(report, read_only=True, data_only=True)
    try:
        ws = wb["가이드위반"]
        sections = []
        current = None
        for cells in ws.iter_rows(min_col=2, max_col=11, values_only=True):
            number, carrier, plan, monthly, period, after = cells[:6]
            rule = cells[9]
            if isinstance(number, str) and "가이드 위반" in number:
                current = {"title": number, "rule": rule or "", "rows": []}
                sections.append(current)
            elif current is not None and isinstance(number, int) and carrier:
                current["rows"].append((number, carrier, plan, monthly, period, after))
        if len(sections) != 2:
            raise RuntimeError("가이드위반 시트에서 RS/RM 구역을 찾지 못했습니다")
        return sections
    finally:
        wb.close()


def email_body(sections, report):
    title = html.escape(report.stem)
    blocks = [f"<h2>MVNO 요금제 가이드 위반 — {title}</h2>"]
    for section in sections:
        rows = section["rows"]
        blocks.append(f"<h3>{html.escape(section['title'])}: {len(rows)}건</h3>")
        blocks.append(f"<p style='white-space:pre-line'>{html.escape(section['rule'])}</p>")
        blocks.append("<table border='1' cellspacing='0' cellpadding='5'><thead><tr>"
                      + "".join(f"<th>{html.escape(s)}</th>" for s in
                                ("순번", "사업자", "요금제 명", "월 요금", "할인 기간", "기간 이후 요금"))
                      + "</tr></thead><tbody>")
        for row in rows:
            blocks.append("<tr>" + "".join(f"<td>{html.escape(str(v if v is not None else ''))}</td>"
                                             for v in row) + "</tr>")
        blocks.append("</tbody></table>")
    blocks.append("<p>전체 요금제 및 전일 대비 변동은 첨부 엑셀을 확인하세요.</p>")
    return "\n".join(blocks)


def telegram_messages(sections, report):
    header = f"MVNO 가이드위반 {report.stem}\n" + " / ".join(
        f"{'RS' if ' RS ' in section['title'] else 'RM'} {len(section['rows'])}건"
        for section in sections)
    lines = [header]
    for section in sections:
        label = "RS" if " RS " in section["title"] else "RM"
        lines.append(f"\n[{label}]")
        for _, carrier, plan, monthly, period, after in section["rows"]:
            lines.append(f"• {carrier} | {plan} | {monthly}원 / {period if period is not None else '-'}개월 / 이후 {after if after is not None else '-'}원")
    # 한 항목이 아주 길더라도 텔레그램의 메시지 제한을 넘지 않도록 분할.
    messages, chunk = [], ""
    for line in lines:
        for part in (line[i:i + 3000] for i in range(0, len(line), 3000)):
            if len(chunk) + len(part) + 1 > 3500:
                messages.append(chunk)
                chunk = ""
            chunk += ("\n" if chunk else "") + part
    if chunk:
        messages.append(chunk)
    return messages


def required(*names):
    missing = [name for name in names if not os.environ.get(name)]
    if missing:
        raise RuntimeError("GitHub Actions Secret 누락: " + ", ".join(missing))


def send_mail(report, sections):
    required("GMAIL_ADDRESS", "GMAIL_APP_PASSWORD", "MVNO_MAIL_TO")
    sender = os.environ["GMAIL_ADDRESS"]
    message = EmailMessage()
    message["From"] = sender
    message["To"] = os.environ["MVNO_MAIL_TO"]
    message["Subject"] = f"MVNO 가이드위반 보고 {report.stem}"
    message.set_content("가이드위반 시트의 내용은 HTML 메일 본문 및 첨부 엑셀을 확인하세요.")
    message.add_alternative(email_body(sections, report), subtype="html")
    message.add_attachment(report.read_bytes(), maintype="application",
                           subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           filename=report.name)
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=60) as smtp:
        smtp.login(sender, os.environ["GMAIL_APP_PASSWORD"].replace(" ", ""))
        smtp.send_message(message)
    print("메일 발송 완료")


def send_telegram(report, sections):
    required("MVNO_TELEGRAM_BOT_TOKEN", "MVNO_TELEGRAM_CHAT_ID")
    token = os.environ["MVNO_TELEGRAM_BOT_TOKEN"]
    for message in telegram_messages(sections, report):
        body = urllib.parse.urlencode({"chat_id": os.environ["MVNO_TELEGRAM_CHAT_ID"],
                                      "text": message}).encode("utf-8")
        request = urllib.request.Request(f"https://api.telegram.org/bot{token}/sendMessage",
                                         data=body, method="POST")
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                if not json.load(response).get("ok"):
                    raise RuntimeError("텔레그램 API가 요청을 거부했습니다")
        except Exception as exc:
            # HTTPError 등에 토큰이 포함된 URL이 찍히지 않도록 상세 오류는 숨긴다.
            raise RuntimeError("텔레그램 발송 실패: Bot token, Chat ID, 봇 대화 시작 여부를 확인하세요") from None
    print("텔레그램 발송 완료")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("report", type=Path)
    args = parser.parse_args()
    sections = guide_sections(args.report)
    send_mail(args.report, sections)
    send_telegram(args.report, sections)
