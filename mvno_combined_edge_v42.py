"""알뜰폰 사업자 요금제 통합 크롤러 (Microsoft Edge 전용)."""

import re
import html
import json
import argparse
from collections import Counter, defaultdict
from datetime import date, timedelta
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

FREET_URL = "https://www.freet.co.kr/plan/ratePlan"
MOBING_URL = "https://www.mobing.co.kr/product/plan"
ASIA_URL = "https://asiamobile.kr/view/price/pricePlan.aspx"
SUGAR_URLS = [
    "https://www.sugarmobile.co.kr/rate_plan.do?type=T012",
    "https://www.sugarmobile.co.kr/rate_plan.do?type=T016",
    "https://www.sugarmobile.co.kr/rate_plan.do?type=T006",
    "https://www.sugarmobile.co.kr/rate_plan.do?type=T005",
]
SIWOL_URLS = [
    "https://siwolmobile.com/rate_plan.do?type=T006",
    "https://siwolmobile.com/rate_plan.do?type=T002",
    "https://siwolmobile.com/rate_plan.do?type=T005",
]
KG_URL = "https://www.kgmobile.co.kr/plan/all"
EYAGI_URL = "https://www.eyagi.co.kr/shop/plan/list.php?tag=pick"
EYES_URL = "https://www.eyes.co.kr/payplan/all_plan"
CHANCE_URL = "https://chancemobile.co.kr/view/plan/phone_plan.aspx"
MONA_URL = "https://mobilemona.co.kr/view/plan/rate_plan.aspx"
AJD_URLS = [f"https://www.ajdmobile.co.kr/rate_plan.do?type={type_code}"
            for type_code in ("T010", "T012", "T006", "T011")]
INS_URLS = [f"https://www.insmobile.co.kr/rate_plan.do?type={type_code}"
            for type_code in ("T007", "T003", "T002", "T004", "T011")]
TPLUS_URL = "https://www.tplusmobile.com/main/rate/join"
TEMPLATE_FILE = Path(__file__).with_name("mvno_combined_plans_reference.xlsx")
COLUMNS = ["사업자", "요금제", "통화 제공량", "문자 제공량", "데이터 제공량",
           "망구분", "LTE/5G 구분", "월 요금", "할인 기간", "기간 이후 요금",
           "요금구분", "홈페이지"]


def clean(s):
    return re.sub(r"\s+", " ", s or "").strip()


def num(s):
    m = re.search(r"[0-9][0-9,]*", s or "")
    return m.group(0).replace(",", "") if m else ""


def normalize_units(s):
    """필터링을 위해 숫자와 단위 사이 불필요한 공백을 제거."""
    s = clean(s)
    s = re.sub(r"(?<=\d)\s+(?=(?:분|건|GB|MB|Mbps|원)\b)", "", s)
    return s


def after_price(s):
    m = re.search(r"(\d+)\s*개월(?:차부터|\s*후)\s*([\d,]+)\s*원", s)
    return (m.group(1), m.group(2).replace(",", "")) if m else ("", "")


def edge_path():
    paths = [Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
             Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe")]
    path = next((p for p in paths if p.exists()), None)
    if not path:
        raise RuntimeError("Microsoft Edge 실행 파일을 찾지 못했습니다.")
    return path


def click_more(page):
    """더보기 버튼이 있는 사이트에서 목록이 늘지 않을 때까지 클릭."""
    selectors = ["#moreBtn a", "text=더보기", "a:has-text('더보기')", "button:has-text('더보기')"]
    for selector in selectors:
        loc = page.locator(selector).first
        if not loc.count() or not loc.is_visible():
            continue
        old = -1
        while loc.is_visible():
            count = page.locator(".plan-item, .callplan-list__listbox_v2").count()
            if count == old:
                break
            old = count
            loc.click()
            page.wait_for_timeout(800)
            loc = page.locator(selector).first
        return


def crawl_freet(page):
    page.goto(FREET_URL, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_selector(".plan-item", timeout=30000)
    click_more(page)
    rows = []
    cards = page.locator(".plan-item")
    for i in range(cards.count()):
        card = cards.nth(i)
        text = clean(card.inner_text())
        def first(sel):
            x = card.locator(sel)
            return clean(x.first.inner_text()) if x.count() else ""
        icons = card.locator(".plan-icon-list li")
        months, price_after = after_price(text)
        rows.append({
            "사업자": "프리티", "요금제": first(".plan-top .name"),
            "통화 제공량": clean(icons.nth(0).inner_text()) if icons.count() > 0 else "",
            "문자 제공량": clean(icons.nth(1).inner_text()) if icons.count() > 1 else "",
            "데이터 제공량": first(".plan-title .title"), "망구분": first(".flag-type"),
            "LTE/5G 구분": "5G" if "5G" in text else "LTE",
            "월 요금": num(first(".plan-price .price")), "할인 기간": months,
            "기간 이후 요금": price_after, "요금구분": "후불", "홈페이지": "홈페이지",
        })
    print(f"프리티: {len(rows)}건")
    return rows


def crawl_mobing(page):
    page.goto(MOBING_URL, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_selector(".callplan-list__listbox_v2", timeout=30000)
    # 모빙은 20개 단위의 '더보기 페이지네이션'이다.
    # 버튼을 누르면 Vue가 다음 페이지를 받아 기존 목록에 누적한다.
    rows = []
    previous_card_count = 0
    for page_no in range(1, 13):
        more = page.locator(".page-more__btn")
        page_label = clean(page.locator(".page-more__num").inner_text()) if page.locator(".page-more__num").count() else ""
        cards_now = page.locator(".callplan-list__listbox_v2")
        print(f"모빙 현재 페이지 {page_no}/12 ({page_label}): {cards_now.count()}건")
        current_card_count = cards_now.count()
        # 모빙은 다음 페이지 카드가 기존 목록 뒤에 누적되므로 신규 카드만 추출한다.
        for i in range(previous_card_count, current_card_count):
            card = cards_now.nth(i)
            text = clean(card.inner_text())

            def first(selector):
                loc = card.locator(selector)
                return clean(loc.first.inner_text()) if loc.count() else ""

            months, price_after = after_price(text)
            price_text = first(".price__area .price")
            rows.append({
                "사업자": "모빙",
                "요금제": first(".top__area .name"),
                "통화 제공량": first(".item__list.voice .ttx"),
                "문자 제공량": first(".item__list.sms .ttx"),
                "데이터 제공량": first(".item__list.data .ttx"),
                "망구분": "SKT" if "SKT" in text else ("KT" if "KT" in text else "LGU+"),
                "LTE/5G 구분": "5G" if "5G" in text else "LTE",
                "월 요금": num(price_text),
                "할인 기간": months,
                "기간 이후 요금": price_after,
                "요금구분": "후불",
                "홈페이지": "홈페이지",
            })
        previous_card_count = current_card_count
        if page_no == 12:
            break
        if not more.count():
            raise RuntimeError(f"모빙 {page_no}페이지에서 다음 페이지 화살표를 찾지 못했습니다.")
        # 실제 아래 화살표 아이콘을 직접 클릭한다.
        more.last.scroll_into_view_if_needed()
        arrow = more.last.locator(".i-btn-more")
        if arrow.count():
            arrow.click(force=True)
        else:
            more.last.click(force=True)
        # Vue API 호출 및 다음 20개 렌더링을 기다린다.
        page.wait_for_timeout(500)
        try:
            page.wait_for_function(
                "before => document.querySelectorAll('.callplan-list__listbox_v2').length > before",
                arg=current_card_count,
                timeout=12000,
            )
        except PlaywrightTimeoutError:
            raise RuntimeError(
                f"모빙 {page_no + 1}페이지 카드 로딩이 완료되지 않았습니다. "
                f"현재 카드 수: {page.locator('.callplan-list__listbox_v2').count()}"
            )
        new_label = clean(page.locator(".page-more__num").inner_text()) if page.locator(".page-more__num").count() else ""
        if new_label == page_label:
            page.wait_for_timeout(1200)
            new_label = clean(page.locator(".page-more__num").inner_text()) if page.locator(".page-more__num").count() else ""
        if new_label == page_label:
            raise RuntimeError(f"모빙 {page_no}페이지의 아래 화살표 클릭 후 페이지가 바뀌지 않았습니다. 현재 표시: {page_label}")
    print(f"모빙: {len(rows)}건")
    return rows


def crawl_asia(page):
    page.goto(ASIA_URL, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_selector(".plan", timeout=30000)
    rows = []
    # 선불은 제외하고 후불요금제만 수집한다.
    for pay_button in ["#btnPostpay"]:
        button = page.locator(pay_button)
        if button.count():
            button.click()
            page.wait_for_timeout(900)
        all_button = page.locator("#btnAll")
        if all_button.count():
            all_button.click()
            page.wait_for_timeout(900)
        cards = page.locator("#planList .plan")
        for i in range(cards.count()):
            card = cards.nth(i)
            text = clean(card.inner_text())
            def first(selector):
                loc = card.locator(selector)
                return clean(loc.first.inner_text()) if loc.count() else ""
            def text_value(pattern):
                match = re.search(pattern, text, re.IGNORECASE)
                return clean(match.group(1)) if match else ""
            months, price_after = after_price(text)
            pay_type = first(".tag_postpay") or "후불"
            network = first(".plan_tag p")
            # 아시아모바일은 아이콘 옆의 plan_summary 값으로도 제공량을 표시한다.
            voice = first(".plan_detail_call") or first(".plan_summary .plan_call") or text_value(r"(?:음성|통화)\s*([^\n]+)")
            sms = first(".plan_detail_sms") or first(".plan_summary .plan_sms") or text_value(r"문자\s*([^\n]+)")
            data = first(".plan_detail_data") or first(".plan_summary .plan_data") or text_value(r"데이터\s*([^\n]+)")
            rows.append({
                "사업자": "아시아모바일",
                "요금제": first(".plan_tit"),
                "통화 제공량": voice,
                "문자 제공량": sms,
                "데이터 제공량": data,
                "망구분": network,
                "LTE/5G 구분": "5G" if "5G" in text else "LTE",
                "월 요금": num(first(".price_now")),
                "할인 기간": months,
                "기간 이후 요금": price_after,
                "요금구분": pay_type,
                "홈페이지": "홈페이지",
            })
    print(f"아시아모바일: {len(rows)}건")
    return rows


def crawl_sugar(page):
    rows = []
    for url in SUGAR_URLS:
        page.goto(url, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_selector(".card_list_item", timeout=30000)
        cards = page.locator(".card_list_item")
        for i in range(cards.count()):
            card = cards.nth(i)
            text = clean(card.inner_text())
            def first(selector):
                loc = card.locator(selector)
                return clean(loc.first.inner_text()) if loc.count() else ""
            title = first(".title")
            network = first(".title .text_light_color")
            plan_name = clean(re.sub(r"^" + re.escape(network), "", title).strip()) if network else title
            badge = first(".badge")
            months_match = re.search(r"(\d+)개월", badge + " " + text)
            months = months_match.group(1) if months_match else ""
            monthly_match = re.search(r"월\s*([\d,]+)\s*원", text)
            later_match = re.search(r"(?:이후|후)\s*([\d,]+)\s*원", text)
            voice = normalize_units(re.sub(r"^음성통화\s*", "", first(".desc li:nth-child(2)")))
            sms = normalize_units(re.sub(r"^문자\s*", "", first(".desc li:nth-child(3)")))
            data = normalize_units(re.sub(r"^월\s*데이터\s*", "", first(".desc li:nth-child(1)")))
            rows.append({
                "사업자": "슈가모바일",
                "요금제": plan_name,
                "통화 제공량": voice,
                "문자 제공량": sms,
                "데이터 제공량": data,
                "망구분": network,
                "LTE/5G 구분": "5G" if "5G" in text or "T005" in url or "T016" in url else "LTE",
                "월 요금": monthly_match.group(1).replace(",", "") if monthly_match else "",
                "할인 기간": months,
                "기간 이후 요금": later_match.group(1).replace(",", "") if later_match else "",
                "요금구분": "후불",
                "홈페이지": "홈페이지",
            })
        print(f"슈가모바일 {url.split('=')[-1]}: {cards.count()}건")
    print(f"슈가모바일 합계: {len(rows)}건")
    return rows


def crawl_siwol(page):
    rows = []
    for url in SIWOL_URLS:
        page.goto(url, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_selector(".rate_list .card_list_item", timeout=30000)
        cards = page.locator(".rate_list .card_list_item")
        for i in range(cards.count()):
            card = cards.nth(i)

            def first(selector):
                loc = card.locator(selector)
                return clean(loc.first.inner_text()) if loc.count() else ""

            title = first(".title")
            network = first(".title .text_light_color")
            plan_name = clean(title.removeprefix(network)) if network else title
            data = normalize_units(first(".desc .ico_data"))
            voice = normalize_units(re.sub(r"^통화\s*", "", first(".desc .ico_voice")))
            sms = normalize_units(re.sub(r"^문자\s*", "", first(".desc .ico_mms")))
            monthly = num(first(".price > span"))
            later = first(".price .ref")
            match = re.search(r"(\d+)\s*개월\s*이후\s*([\d,]+)\s*원", later)
            if not plan_name or not monthly or not data:
                raise RuntimeError(f"시월모바일 {url}의 {i + 1}번째 카드 필수 항목 누락")
            rows.append({
                "사업자": "시월모바일", "요금제": plan_name,
                "통화 제공량": voice, "문자 제공량": sms,
                "데이터 제공량": data, "망구분": network,
                "LTE/5G 구분": "5G" if "5G" in plan_name or "T005" in url else "LTE",
                "월 요금": monthly,
                "할인 기간": match.group(1) if match else "",
                "기간 이후 요금": match.group(2).replace(",", "") if match else "",
                "요금구분": "후불", "홈페이지": "홈페이지",
            })
        print(f"시월모바일 {url.split('=')[-1]}: {cards.count()}건")
    print(f"시월모바일 합계: {len(rows)}건")
    return rows


def kg_text(value):
    """상세 설명의 HTML을 공백 구분 텍스트로 변환한다."""
    return clean(html.unescape(re.sub(r"<[^>]+>", " ", value or "")))


def kg_allowance(value, unit):
    value = str(value or "").strip()
    if value == "-1":
        return "기본제공"
    if value in ("", "-"):
        return ""
    return normalize_units(f"{value}{unit}")


def kg_row(plan):
    description = kg_text(plan.get("contents"))
    lifetime = "할인기간 제한이 없는 평생 할인 요금제" in description
    period = re.search(r"기본료 할인 프로모션은\s*가입월\s*포함\s*(\d+)\s*개월간", description)
    sales = plan.get("saleList") or []
    if len(sales) != 1:
        raise ValueError(f"KG {plan.get('planNo')}: 할인 항목 {len(sales)}개 (가격 확인 필요)")
    sale = sales[0]
    basic = int(plan["basicAmount"])
    forever_discount = int(sale.get("ltSaleAmount") or 0)
    limited_discount = int(sale.get("prSaleAmount") or 0)
    if limited_discount and not period:
        raise ValueError(f"KG {plan.get('planNo')}: 기간 할인 문구를 찾지 못했습니다")
    if lifetime and period:
        raise ValueError(f"KG {plan.get('planNo')}: 평생/기간 할인 문구가 동시에 발견됨")
    monthly = basic - forever_discount - limited_discount
    after = basic - forever_discount
    if monthly < 0 or after < 0:
        raise ValueError(f"KG {plan.get('planNo')}: 할인 후 요금이 음수입니다")

    if plan.get("basicMonthDataInfinite") == 0:
        data = f"매일 {kg_allowance(plan.get('basicDayData'), plan.get('basicDayDataUnit') or 'GB')}"
    else:
        data = kg_allowance(plan.get("basicMonthData"), plan.get("basicMonthDataUnit") or "GB")
        data = f"월 {data}"
    if plan.get("exhaustion"):
        data += f" (소진시 {clean(plan['exhaustion'])})"

    network = {"LGT": "LG U+", "KT": "KT", "SKT": "SKT"}.get(plan.get("telco"), "")
    return {
        "사업자": "KG모바일", "요금제": clean(plan.get("planName")),
        "통화 제공량": kg_allowance(plan.get("basicVoice"), "분"),
        "문자 제공량": kg_allowance(plan.get("basicSms"), "건"),
        "데이터 제공량": normalize_units(data),
        "망구분": network,
        "LTE/5G 구분": "5G" if plan.get("network") == "5G" else "LTE",
        "월 요금": str(monthly),
        "할인 기간": period.group(1) if period and not lifetime else "",
        "기간 이후 요금": str(after) if period and not lifetime else "",
        "요금구분": "후불", "홈페이지": "홈페이지",
    }


def kg_browser_json(page, path):
    """Windows 인증서를 사용하는 Edge 탭에서 동일 출처 API를 호출한다."""
    result = page.evaluate("""async (path) => {
        const controller = new AbortController();
        const timer = setTimeout(() => controller.abort(), 30000);
        try {
            const response = await fetch(path, {
                credentials: 'same-origin', signal: controller.signal,
                headers: { 'Accept': 'application/json' }
            });
            if (!response.ok) throw new Error(`HTTP ${response.status}: ${path}`);
            return await response.json();
        } finally {
            clearTimeout(timer);
        }
    }""", path)
    if result.get("code") != 1:
        raise RuntimeError(f"KG API 응답 오류: {path}, {result.get('message')}")
    return result


def crawl_kg(page):
    # 화면의 1, 2, 3... 페이지는 이 목록 API의 결과를 15개씩 잘라 보여준다.
    # 각 요금제 상세 페이지가 호출하는 상세 API를 개별 조회한다.
    page.goto(KG_URL, wait_until="domcontentloaded", timeout=60000)
    payload = kg_browser_json(page,
        "/api/product/plan?page=1&limit=-1&block=5&isUser=true&useFlag=Y")
    entity = payload.get("entity") or {}
    plans = entity.get("list") or []
    expected = (entity.get("pageInfo") or {}).get("totalRow")
    if not plans or expected is None or len(plans) != int(expected):
        raise RuntimeError(f"KG 목록 건수 불일치: 조회 {len(plans)}, 전체 {expected}")
    print(f"KG모바일 목록: {len(plans)}건 (상세 페이지별 조회 시작)")

    rows = []
    for index, summary in enumerate(plans, 1):
        plan_no = summary["planNo"]
        detail_payload = kg_browser_json(page, f"/api/product/plan/{plan_no}")
        if not detail_payload.get("entity"):
            raise RuntimeError(f"KG 상세 응답 오류: {plan_no}")
        detail = detail_payload["entity"]
        if detail.get("planNo") != plan_no:
            raise RuntimeError(f"KG 상세 요금제 번호 불일치: {plan_no}")
        rows.append(kg_row(detail))
        if index % 15 == 0 or index == len(plans):
            print(f"KG모바일 상세 추출: {index}/{len(plans)}건")
    return rows


def eyagi_rows(payload):
    if payload.get("result") != "SUCCESS":
        raise RuntimeError(f"이야기모바일 목록 응답 오류: {payload.get('result')}")
    data = payload.get("data") or {}
    count = int(payload.get("num") or 0)
    required = ("comm_name", "pay_type", "mno_gubun", "is_5g_yn",
                "free_data_str", "add_data", "qos", "call_str", "sms_str",
                "selling_price", "terms", "restoration_fee")
    if count == 0 or any(len(data.get(key) or []) != count for key in required):
        raise RuntimeError("이야기모바일 목록 필드 또는 전체 건수가 일치하지 않습니다")

    rows = []
    for i in range(count):
        # 사이트 응답의 1=선불, 2=후불. 선불과 유형 미지정은 제외한다.
        if str(data["pay_type"][i]) != "2":
            continue
        plan_name = clean(data["comm_name"][i])
        amount = num(str(data["selling_price"][i]))
        base_data = clean(data["free_data_str"][i])
        extra = clean(data["add_data"][i])
        qos = clean(str(data["qos"][i] or ""))
        data_parts = [base_data]
        if extra:
            data_parts.append(extra)
        allowance = "월 " + " + ".join(part for part in data_parts if part)
        if qos and qos not in ("999", "0"):
            allowance += f" (소진시 {qos}Mbps)" if qos.isdigit() else f" (소진시 {qos})"
        term = clean(str(data["terms"][i] or ""))
        after = num(str(data["restoration_fee"][i] or "")) if term not in ("", "99") else ""
        if not plan_name or not amount or not base_data or (term not in ("", "99") and not after):
            raise RuntimeError(f"이야기모바일 {i + 1}번째 요금제 필수 값 누락: {plan_name}")
        network = {"lgt": "LG U+", "skt": "SKT", "kt": "KT"}.get(
            str(data["mno_gubun"][i]).lower(), "")
        rows.append({
            "사업자": "큰사람", "요금제": plan_name,
            "통화 제공량": normalize_units(clean(data["call_str"][i])),
            "문자 제공량": normalize_units(clean(data["sms_str"][i])),
            "데이터 제공량": normalize_units(allowance),
            "망구분": network,
            "LTE/5G 구분": "5G" if clean(data["is_5g_yn"][i]).upper() == "5G" else "LTE",
            "월 요금": amount,
            "할인 기간": term if term not in ("", "99") else "",
            "기간 이후 요금": after,
            "요금구분": "후불", "홈페이지": "홈페이지",
        })
    return rows


def crawl_eyagi(page):
    page.goto(EYAGI_URL, wait_until="domcontentloaded", timeout=60000)
    payload = page.evaluate("""async () => {
        const form = new URLSearchParams({page_seq: '2', mode: 'comm_list_view'});
        const controller = new AbortController();
        const timer = setTimeout(() => controller.abort(), 60000);
        try {
            const response = await fetch('/shop/plan/json_proc.php', {
                method: 'POST', credentials: 'same-origin',
                headers: {'Content-Type': 'application/x-www-form-urlencoded'},
                body: form.toString(), signal: controller.signal
            });
            if (!response.ok) throw new Error(`HTTP ${response.status}`);
            return await response.json();
        } finally { clearTimeout(timer); }
    }""")
    rows = eyagi_rows(payload)
    print(f"큰사람: 후불 {len(rows)}건 (전체 {payload['num']}건)")
    return rows


def eyes_card(card):
    def first(selector):
        loc = card.locator(selector)
        return clean(loc.first.inner_text()) if loc.count() else ""

    name = first(".mid > div:first-child > p.body_medium.sb")
    data = first(".plan_tit")
    info = card.locator(".plan_info > p")
    voice = clean(info.nth(0).inner_text()) if info.count() > 0 else ""
    sms = clean(info.nth(1).inner_text()) if info.count() > 1 else ""
    network_label = first(".label_group .badge.skt, .label_group .badge.kt, .label_group .badge.lg")
    network = {"LGU+": "LG U+", "SKT": "SKT", "KT": "KT"}.get(network_label.upper(), "")
    monthly = num(first(".price .current_p"))
    period_text = first(".price .period")
    months = re.search(r"(?:첫\s*)?(\d+)\s*개월", period_text)
    period = months.group(1) if months else ""
    after = ""
    tooltip = card.locator(".price .tooltip_btn[data-discount]")
    if tooltip.count():
        stages = json.loads(tooltip.first.get_attribute("data-discount"))
        if period:
            if len(stages) < 2 or not isinstance(stages[1].get("price"), (int, float)):
                raise RuntimeError(f"아이즈모바일 {name}: 할인 이후 요금 단계 누락")
            after = str(int(stages[1]["price"]))
    if period and not after:
        after = num(first(".price .org_p"))
    if not name or not data or not monthly or not network or info.count() < 2 or (period and not after):
        raise RuntimeError(f"아이즈모바일 카드 필수 값 누락: {name} / {data} / {network}")
    if re.search(r"선불|prepaid", name, re.I):
        return None
    # 요금제명의 5G 또는 상세 URL 분류 코드가 명시된 경우를 사용한다.
    tech = "5G" if re.search(r"5G", name, re.I) or "/C05" in (card.get_attribute("href") or "") else "LTE"
    return {
        "사업자": "아이즈비전", "요금제": name,
        "통화 제공량": normalize_units(voice), "문자 제공량": normalize_units(sms),
        "데이터 제공량": normalize_units("월 " + re.sub(r"^월\s*", "", data)),
        "망구분": network, "LTE/5G 구분": tech,
        "월 요금": monthly, "할인 기간": period, "기간 이후 요금": after,
        "요금구분": "후불", "홈페이지": "홈페이지",
    }


def crawl_eyes(page, output_rows):
    page.goto(EYES_URL, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_selector(".plancard_list a.plan_card.full", timeout=30000)
    expected = num(page.locator(".filter_apply_box .tot_cnt").first.inner_text())
    if not expected:
        raise RuntimeError("아이즈비전 전체 요금제 수를 읽지 못했습니다")
    total_pages = (int(expected) + 19) // 20
    seen, count = set(), 0
    for page_no in range(1, total_pages + 1):
        if page_no > 1:
            # 사이트 페이지 버튼의 onclick에 있는 실제 주소 형식이다.
            page.goto(f"https://www.eyes.co.kr//payplan/all_plan//?&page={page_no}",
                      wait_until="domcontentloaded", timeout=60000)
            page.wait_for_selector(".plancard_list a.plan_card.full", timeout=30000)
        active = clean(page.locator(".list_footer.paging button.pagenation.active").first.inner_text())
        if active != str(page_no):
            raise RuntimeError(f"아이즈비전 요청 {page_no}페이지, 실제 표시 {active}페이지")
        cards = page.locator(".plancard_list a.plan_card.full")
        first_href = cards.first.get_attribute("href")
        if first_href in seen:
            raise RuntimeError(f"아이즈비전 {page_no}페이지: 동일한 첫 요금제 반복")
        seen.add(first_href)
        page_rows = [row for card in cards.all() if (row := eyes_card(card)) is not None]
        if not page_rows:
            raise RuntimeError(f"아이즈비전 {page_no}페이지: 후불 요금제 행 없음")
        output_rows.extend(page_rows)
        count += len(page_rows)
        print(f"아이즈비전 {page_no}페이지: {len(page_rows)}건 (누적 {count}건)")
    if expected and count != int(expected):
        raise RuntimeError(f"아이즈비전 건수 불일치: 수집 {count}건 / 화면 전체 {expected}건")
    print(f"아이즈비전 합계: {count}건")


def chance_rows(payload):
    if payload.get("RESULT") != "Y" or not isinstance(payload.get("DATA"), list):
        raise RuntimeError(f"찬스모바일 목록 응답 오류: {payload.get('RESULTMSG')}")
    plans = payload["DATA"]
    if not plans:
        raise RuntimeError("찬스모바일 목록이 비어 있습니다")
    rows = []
    for i, plan in enumerate(plans, 1):
        name = clean(plan.get("GDNM"))
        if re.search(r"선불|prepaid", name, re.I):
            continue
        period = num(clean(plan.get("LIMPERIOD")))
        base = num(str(plan.get("TT_AMT") or ""))
        discount = num(str(plan.get("DISCOUNT") or ""))
        monthly = discount if discount != "0" else base
        later = num(str(plan.get("ORGCHARGE") or "")) if period else ""
        network = {"LGU+": "LG U+", "KT": "KT", "SKT": "SKT"}.get(clean(plan.get("MNO_CD")).upper(), "")
        tech = clean(plan.get("NETDIV")).upper()
        base_data = clean(plan.get("DATAAMOUNT"))
        if not name or not base or monthly == "" or not base_data or not network or tech not in ("LTE", "5G") or (period and not later):
            raise RuntimeError(f"찬스모바일 {i}번째 요금제 필수 값 누락: {name}")
        data = base_data + clean(plan.get("ADD_DATA"))
        if clean(plan.get("QOSFG")) == "1":
            qos = clean(plan.get("QOSAMT"))
            if qos and qos != "0":
                data += f" (소진시 {qos.lstrip('+ ').strip()})"
        if clean(plan.get("BLOCKYN")) == "Y":
            data += " (소진후 데이터 차단)"
        voice = clean(plan.get("VOICEAMOUNT"))
        extra = clean(plan.get("VOICE_ADD_AMT"))
        if extra and extra != "0":
            voice += f" (부가 음성 {extra}분)"
        rows.append({
            "사업자": "찬스모바일", "요금제": name,
            "통화 제공량": normalize_units(voice),
            "문자 제공량": normalize_units(clean(plan.get("LETTERAMOUNT"))),
            "데이터 제공량": normalize_units("월 " + data),
            "망구분": network, "LTE/5G 구분": tech,
            "월 요금": monthly, "할인 기간": period, "기간 이후 요금": later,
            "요금구분": "후불", "홈페이지": "홈페이지",
        })
    return rows


def crawl_chance(page):
    page.goto(CHANCE_URL, wait_until="domcontentloaded", timeout=60000)
    # 페이지 자체에서 사용하는 AjaxPhone_plan.aspx의 type 02(전체 목록) 요청.
    payload = page.evaluate("""async () => {
        const order = document.querySelector('#selOrder')?.value || '';
        const [order_type = '', order_align = ''] = order.split('_');
        const body = JSON.stringify({
            header: [{type: '02'}],
            body: [{seq: '', order_type, order_align}]
        });
        const response = await fetch('/common/component/plan/AjaxPhone_plan.aspx', {
            method: 'POST', credentials: 'same-origin',
            headers: {'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
                      'X-Requested-With': 'XMLHttpRequest'},
            body
        });
        if (!response.ok) throw new Error(`찬스모바일 HTTP ${response.status}`);
        return await response.json();
    }""")
    rows = chance_rows(payload)
    print(f"찬스모바일: {len(rows)}건 (전체 {len(payload['DATA'])}건)")
    return rows


def mona_card(card):
    def first(selector):
        loc = card.locator(selector)
        return clean(loc.first.inner_text()) if loc.count() else ""

    name = first(".pb-plan-item_name")
    if re.search(r"선불|prepaid", name, re.I):
        return None
    tech = card.locator(".netdiv")
    tech_value = clean(tech.first.get_attribute("data-value")) if tech.count() else ""
    data = first(".pb-plan-data_name.data")
    voice = first(".pb-plan-data_name.voice")
    sms = first(".pb-plan-data_name.letter")
    price = card.locator(".discount")
    monthly = num(price.first.get_attribute("data-value") or price.first.inner_text()) if price.count() else ""
    period = card.locator(".event-period")
    period_label = first(".event-period")
    period_value = num(period.first.get_attribute("data-value") or period_label) if period.count() and re.search(r"\d+\s*개월", period_label) else ""
    later = num(first(".pb-is-linethrough")) if period_value else ""
    if not name or not data or not voice or not sms or not monthly or tech_value not in ("LTE", "5G") or (period_value and not later):
        raise RuntimeError(f"코나아이 요금제 필수 값 누락: {name}")
    return {
        "사업자": "코나아이", "요금제": name,
        "통화 제공량": normalize_units(voice), "문자 제공량": normalize_units(sms),
        "데이터 제공량": normalize_units("월 " + re.sub(r"^월\s*", "", data)),
        "망구분": "LG U+", "LTE/5G 구분": tech_value,
        "월 요금": monthly, "할인 기간": period_value, "기간 이후 요금": later,
        "요금구분": "후불", "홈페이지": "홈페이지",
    }


def crawl_mona(page):
    page.goto(MONA_URL, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_function("document.querySelectorAll('.pb-plan-select_list > li.pb-plan-item').length > 0", timeout=45000)
    cards = page.locator(".pb-plan-select_list > li.pb-plan-item")
    displayed = num(page.locator("#totalCnt").inner_text())
    if not displayed or cards.count() != int(displayed):
        raise RuntimeError(f"코나아이 목록 수 불일치: 카드 {cards.count()} / 화면 {displayed}")
    rows = [row for card in cards.all() if (row := mona_card(card)) is not None]
    print(f"코나아이: 후불 {len(rows)}건 (전체 {displayed}건)")
    return rows


def crawl_ajd(page):
    rows = []
    for url in AJD_URLS:
        page.goto(url, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_selector("#postLTE .card_list_item", timeout=30000)
        cards = page.locator("#postLTE .card_list_item")
        page_rows = []
        for i in range(cards.count()):
            card = cards.nth(i)
            def first(selector):
                loc = card.locator(selector)
                return clean(loc.first.inner_text()) if loc.count() else ""

            name = first(".title")
            if re.search(r"선불|prepaid", name, re.I):
                continue
            network_text = first(".badge_wrap > .badge:first-child").upper()
            network = {"LGU+": "LG U+", "KT": "KT", "SKT": "SKT"}.get(network_text, "")
            badges = card.locator(".badge_wrap").first.inner_text()
            desc = card.locator(".detail .desc > li")
            data = clean(desc.nth(0).inner_text()) if desc.count() > 0 else ""
            voice = clean(desc.nth(1).inner_text()) if desc.count() > 1 else ""
            sms = clean(desc.nth(2).inner_text()) if desc.count() > 2 else ""
            voice = re.sub(r"^통화\s*", "", voice)
            sms = re.sub(r"^문자\s*", "", sms)
            price = card.locator(".price").first
            monthly = num(clean(price.evaluate("el => Array.from(el.childNodes).filter(n => n.nodeType === 3).map(n => n.textContent).join(' ')")))
            later_text = first(".price .ref")
            discount = re.search(r"(\d+)\s*개월\s*이후\s*([\d,]+)\s*원", later_text)
            if not name or not network or not data or monthly == "" or desc.count() < 3 or (later_text and not discount):
                raise RuntimeError(f"아정당 {url.split('=')[-1]} {i+1}번째 요금제 필수 값 누락: {name}")
            page_rows.append({
                "사업자": "아정당", "요금제": name,
                "통화 제공량": normalize_units(voice),
                "문자 제공량": normalize_units(sms),
                "데이터 제공량": normalize_units("월 " + re.sub(r"^월\s*", "", data)),
                "망구분": network,
                "LTE/5G 구분": "5G" if "5G" in badges or re.search(r"(?<!\w)5G(?!\w)", name, re.I) else "LTE",
                "월 요금": monthly,
                "할인 기간": discount.group(1) if discount else "",
                "기간 이후 요금": discount.group(2).replace(",", "") if discount else "",
                "요금구분": "후불", "홈페이지": "홈페이지",
            })
        rows.extend(page_rows)
        print(f"아정당 {url.split('=')[-1]}: {len(page_rows)}건")
    print(f"아정당 합계: {len(rows)}건")
    return rows


def ins_detail_prices(page, links):
    """목록에 없는 할인 종료 가격을 같은 출처의 상세 HTML에서 조회."""
    return page.evaluate("""async (links) => {
        const results = [];
        for (let i = 0; i < links.length; i += 5) {
            const batch = links.slice(i, i + 5);
            const values = await Promise.all(batch.map(async (href) => {
                const controller = new AbortController();
                const timer = setTimeout(() => controller.abort(), 20000);
                try {
                    const r = await fetch(href, {credentials: 'same-origin', signal: controller.signal});
                    if (!r.ok) throw new Error(`HTTP ${r.status}: ${href}`);
                    const doc = new DOMParser().parseFromString(await r.text(), 'text/html');
                    const label = doc.querySelector('.rate_info .price .ref')?.textContent?.trim() || '';
                    if (!doc.querySelector('.rate_info .header')) throw new Error(`상세 페이지 누락: ${href}`);
                    return label;
                } finally { clearTimeout(timer); }
            }));
            results.push(...values);
        }
        return results;
    }""", links)


def crawl_ins(page, output_rows):
    total = 0
    for url in INS_URLS:
        page.goto(url, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_selector("#postLTE .card_list_item", timeout=30000)
        cards = page.locator("#postLTE .card_list_item")
        page_rows, detail_links, detail_indices = [], [], []
        for i in range(cards.count()):
            card = cards.nth(i)
            def first(selector):
                loc = card.locator(selector)
                return clean(loc.first.inner_text()) if loc.count() else ""

            title = first(".title")
            if re.search(r"선불|prepaid", title, re.I):
                continue
            network_label = first(".title .text_light_color")
            network = {"LGU+": "LG U+", "KT": "KT", "SKT": "SKT"}.get(network_label.upper(), "")
            name = clean(title.removeprefix(network_label)) if network_label else title
            badges = first(".badge_wrap")
            months = re.search(r"(\d+)\s*개월\s*할인", badges)
            period = months.group(1) if months else ""
            desc = card.locator(".detail .desc > li")
            data = clean(desc.nth(0).inner_text()) if desc.count() > 0 else ""
            voice = re.sub(r"^통화\s*", "", clean(desc.nth(1).inner_text())) if desc.count() > 1 else ""
            sms = re.sub(r"^문자\s*", "", clean(desc.nth(2).inner_text())) if desc.count() > 2 else ""
            monthly = num(first(".price"))
            if not name or not network or not data or monthly == "" or desc.count() < 3 or (not period and "평생" not in badges):
                raise RuntimeError(f"인스모바일 {url.split('=')[-1]} {i+1}번째 요금제 필수 값 누락: {name}")
            row = {
                "사업자": "인스모바일", "요금제": name,
                "통화 제공량": normalize_units(voice), "문자 제공량": normalize_units(sms),
                "데이터 제공량": normalize_units("월 " + re.sub(r"^월\s*", "", data)),
                "망구분": network,
                "LTE/5G 구분": "5G" if "5G" in badges or "(5G)" in name else "LTE",
                "월 요금": monthly, "할인 기간": period, "기간 이후 요금": "",
                "요금구분": "후불", "홈페이지": "홈페이지",
            }
            if period:
                link = card.locator("a.card_rate_link").get_attribute("href")
                if not link:
                    raise RuntimeError(f"인스모바일 상세 링크 누락: {name}")
                detail_links.append(link)
                detail_indices.append(len(page_rows))
            page_rows.append(row)
        labels = ins_detail_prices(page, detail_links) if detail_links else []
        for index, label in zip(detail_indices, labels):
            m = re.search(r"(\d+)\s*개월\s*이후\s*([\d,]+)\s*원", label)
            if not m or m.group(1) != page_rows[index]["할인 기간"]:
                raise RuntimeError(f"인스모바일 할인 요금 확인 실패: {page_rows[index]['요금제']} / {label}")
            page_rows[index]["기간 이후 요금"] = m.group(2).replace(",", "")
        output_rows.extend(page_rows)
        total += len(page_rows)
        print(f"인스모바일 {url.split('=')[-1]}: {len(page_rows)}건 (누적 {total}건)")
    print(f"인스모바일 합계: {total}건")


def tplus_card(card):
    def first(selector):
        loc = card.locator(selector)
        return clean(loc.first.inner_text()) if loc.count() else ""

    name = first(".cardBody .titleArea h4.title")
    if re.search(r"선불|prepaid", name, re.I):
        return None
    network_text = first(".cardHead .badgeArea i.badge:first-child").upper()
    network = {"LGU+": "LG U+", "KT": "KT", "SKT": "SKT"}.get(network_text, "")
    badge = first(".cardHead .badgeArea i.badge.custom")
    months = re.search(r"(\d+)\s*개월", badge)
    period = months.group(1) if months else ""
    data = first(".rateFoot .descArea p.desc20px700")
    voice = first(".rateFoot .descArea .ico.call")
    sms = first(".rateFoot .descArea .ico.message")
    monthly = num(first(".amountArea .textPoint"))
    # 접힌 혜택 설명은 inner_text()가 빈 문자열일 수 있다. 항상 표시되는 취소선 기본료를 사용한다.
    later = num(first(".amountArea .throthDesc")) if period else ""
    if period and later == "":
        raise RuntimeError(f"KCT 할인 이후 요금 확인 실패: {name}")
    if not name or not network or not data or not voice or not sms or monthly == "" or (not period and "평생" not in badge):
        raise RuntimeError(f"KCT 요금제 필수 값 누락: {name}")
    # '매일5G'와 같은 데이터 제공량 표기는 5G 통신망을 뜻하지 않는다.
    tech = "5G" if re.match(r"^5G(?:\s|_)", name, re.I) else "LTE"
    return {
        "사업자": "KCT", "요금제": name,
        "통화 제공량": normalize_units(voice), "문자 제공량": normalize_units(sms),
        "데이터 제공량": normalize_units("월 " + re.sub(r"^월\s*", "", data)),
        "망구분": network, "LTE/5G 구분": tech,
        "월 요금": monthly, "할인 기간": period,
        "기간 이후 요금": later,
        "요금구분": "후불", "홈페이지": "홈페이지",
    }


def crawl_tplus(page):
    page.goto(TPLUS_URL, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_selector("#result_area > .cardArea.accordionsType", timeout=45000)
    cards = page.locator("#result_area > .cardArea.accordionsType")
    total = num(page.locator("#moreTotal").inner_text())
    if not total:
        raise RuntimeError("KCT 전체 요금제 수를 읽지 못했습니다")
    for _ in range(40):
        count = cards.count()
        if count >= int(total):
            break
        more = page.locator("#board_paging button.moreView")
        if not more.count() or not more.is_visible():
            raise RuntimeError(f"KCT 더보기 버튼 없음: {count}/{total}건")
        more.click()
        page.wait_for_function(
            "old => document.querySelectorAll('#result_area > .cardArea.accordionsType').length > old",
            arg=count, timeout=20000)
        print(f"KCT 더보기: {cards.count()}/{total}건")
    if cards.count() != int(total):
        raise RuntimeError(f"KCT 전체 건수 불일치: 카드 {cards.count()} / 화면 {total}")
    rows = [row for card in cards.all() if (row := tplus_card(card)) is not None]
    print(f"KCT: 후불 {len(rows)}건 (화면 전체 {total}건)")
    return rows


def numeric(value):
    """빈칸은 None, 정수 금액/개월은 int로 비교한다."""
    if value is None or isinstance(value, bool):
        return None
    s = str(value).strip().replace(",", "").replace("원", "").replace("개월", "")
    return int(s) if re.fullmatch(r"\d+", s) else None


def plan_key(row):
    """가격 세 칸(H:I:J)을 제외한 요금제 정체성."""
    def field(name):
        return clean(str(row.get(name) or ""))
    network = field("망구분").upper().replace(" ", "")
    return (field("사업자"), field("요금제"), network,
            field("LTE/5G 구분"), field("통화 제공량"),
            field("문자 제공량"), field("데이터 제공량"), field("요금구분"))


def price_key(row):
    return tuple(numeric(row.get(name)) for name in ("월 요금", "할인 기간", "기간 이후 요금"))


def compare_daily(previous, current):
    """동일 이름의 여러 요금제도 삭제하지 않고 정확히 일치하는 행부터 짝지음."""
    old_groups, new_groups = defaultdict(Counter), defaultdict(Counter)
    for row in previous:
        old_groups[plan_key(row)][price_key(row)] += 1
    for row in current:
        new_groups[plan_key(row)][price_key(row)] += 1
    changes = []
    for key in sorted(set(old_groups) | set(new_groups)):
        old, new = old_groups[key], new_groups[key]
        for price in set(old) & set(new):
            both = min(old[price], new[price])
            old[price] -= both
            new[price] -= both
        old_rest = sorted((p for p, qty in old.items() for _ in range(qty)), key=str)
        new_rest = sorted((p for p, qty in new.items() for _ in range(qty)), key=str)
        paired = min(len(old_rest), len(new_rest))
        for idx in range(paired):
            before, after = old_rest[idx], new_rest[idx]
            kind = ", ".join(label for i, label in enumerate(("월요금변경", "할인기간변경", "이후요금변경"))
                             if before[i] != after[i])
            changes.append((kind, *key[:2], key[2], key[3],
                            before[0], after[0], before[1], after[1], before[2], after[2]))
        for before in old_rest[paired:]:
            changes.append(("삭제", *key[:2], key[2], key[3], before[0], None, before[1], None, before[2], None))
        for after in new_rest[paired:]:
            changes.append(("추가", *key[:2], key[2], key[3], None, after[0], None, after[1], None, after[2]))
    return changes


def sheet_rows(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["홈페이지"] if "홈페이지" in book else book.active
        if tuple(sheet.cell(1, i).value for i in range(1, 13)) != tuple(COLUMNS):
            raise RuntimeError(f"{path}: 요금제 열 A:L 형식이 다릅니다")
        return [dict(zip(COLUMNS, values[:12])) for values in sheet.iter_rows(min_row=2, values_only=True)
                if values[0] is not None]
    finally:
        book.close()


def rule_for(row, rs):
    if clean(str(row.get("망구분") or "")).upper().replace(" ", "") not in ("LGU+", "LGT"):
        return ""
    period = numeric(row.get("할인 기간"))
    monthly = numeric(row.get("월 요금"))
    if rs:
        return "RS 할인기간 6개월 이하" if period is not None and period <= 6 else ""
    reasons = []
    if period is not None and period <= 5:
        reasons.append("RM 할인기간 5개월 이하")
    if monthly == 0:
        reasons.append("RM 월요금 0원")
    return ", ".join(reasons)


def write_report(rows, template, previous, output):
    book = load_workbook(template)
    required = {"홈페이지", "참조", "가이드위반", "전일 대비 변동"}
    if not required.issubset(book.sheetnames):
        raise RuntimeError(f"템플릿 필수 시트 누락: {required - set(book.sheetnames)}")
    # Excel의 B열(=C&D&E)과 같은 방식으로 키를 구성한다. 수식 캐시가 없어도 동작한다.
    reference = {
        "".join(str(value or "") for value in (voice, message, data)): str(rs or "").strip()
        for _, _, voice, message, data, _, rs in book["참조"].iter_rows(min_row=3, max_col=7, values_only=True)
        if any(value is not None for value in (voice, message, data))
    }
    ws = book["홈페이지"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for col, label in enumerate([*COLUMNS, "RS구분", "가이드위반"], 1):
        ws.cell(1, col, label)
    rs_violations, rm_violations = [], []
    highlight = PatternFill(fill_type="solid", fgColor="FFF2CC")
    for index, row in enumerate(rows, 2):
        for col, label in enumerate(COLUMNS, 1):
            value = row.get(label)
            ws.cell(index, col, numeric(value) if col in (8, 9, 10) and numeric(value) is not None else (value or None))
        key = "".join(str(row.get(label) or "") for label in
                      ("통화 제공량", "문자 제공량", "데이터 제공량"))
        is_rs = reference.get(key) == "RS"
        ws.cell(index, 13,
                f'=IFERROR(VLOOKUP($C{index}&$D{index}&$E{index},참조!$B:$G,6,0),"")')
        violation = rule_for(row, is_rs)
        ws.cell(index, 14, violation)
        if violation:
            ws.cell(index, 14).fill = highlight
            (rs_violations if is_rs else rm_violations).append(row)
    ws.auto_filter.ref = f"A1:N{max(1, len(rows) + 1)}"
    ws.freeze_panes = "A2"
    ws.column_dimensions["N"].width = 32

    guide = book["가이드위반"]
    rs_title, rs_rule = guide["B2"].value, guide["K2"].value
    rm_title, rm_rule = guide["B18"].value, guide["K18"].value
    labels = [guide.cell(3, col).value for col in range(2, 8)]
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.value = None
    guide["B2"] = rs_title
    guide["K2"] = rs_rule
    for col, label in enumerate(labels, 2):
        guide.cell(3, col, label)

    def put_violations(start, items):
        for offset, item in enumerate(items):
            for col, value in enumerate((offset + 1, item.get("사업자"), item.get("요금제"),
                                         numeric(item.get("월 요금")), numeric(item.get("할인 기간")),
                                         numeric(item.get("기간 이후 요금"))), 2):
                guide.cell(start + offset, col, value)

    put_violations(4, rs_violations)
    rm_header = max(18, 5 + len(rs_violations))
    guide.cell(rm_header, 2, rm_title)
    guide.cell(rm_header, 11, rm_rule)
    for col, label in enumerate(labels, 2):
        guide.cell(rm_header + 1, col, label)
    put_violations(rm_header + 2, rm_violations)

    diff = book["전일 대비 변동"]
    if diff.max_row > 1:
        diff.delete_rows(2, diff.max_row - 1)
    changes = compare_daily(sheet_rows(previous), rows) if previous else []
    for change in changes:
        diff.append(change)
    diff.auto_filter.ref = f"A1:K{max(1, len(changes) + 1)}"
    diff.freeze_panes = "A2"
    book.calculation.fullCalcOnLoad = True
    temp = output.with_name(output.stem + ".tmp.xlsx")
    book.save(temp)
    temp.replace(output)
    print(f"저장 완료: {len(rows)}건 → {output}")
    print(f"가이드위반: RS {len(rs_violations)}건, RM {len(rm_violations)}건")
    print(f"전일 대비: {len(changes)}건" if previous else "전일 파일이 없어 변동 내역은 비워 둡니다.")


def crawl_all():
    rows = []
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True, executable_path=str(edge_path()))
            page = browser.new_page(locale="ko-KR", viewport={"width": 1440, "height": 1200})
            rows.extend(crawl_freet(page))
            rows.extend(crawl_mobing(page))
            rows.extend(crawl_asia(page))
            rows.extend(crawl_sugar(page))
            rows.extend(crawl_siwol(page))
            rows.extend(crawl_eyagi(page))
            rows.extend(crawl_kg(page))
            # 이지모바일은 GitHub 호스팅 러너에서 두 호스트 모두 연결 시간 초과되어 제외.
            crawl_eyes(page, rows)
            rows.extend(crawl_chance(page))
            rows.extend(crawl_mona(page))
            rows.extend(crawl_ajd(page))
            crawl_ins(page, rows)
            rows.extend(crawl_tplus(page))
            browser.close()
    except Exception as exc:
        if rows:
            partial = Path(__file__).with_name(f"mvno_partial_{date.today():%Y%m%d}.xlsx")
            book = Workbook()
            sheet = book.active
            sheet.title = "부분수집"
            sheet.append(COLUMNS)
            for row in rows:
                sheet.append([row.get(label) for label in COLUMNS])
            book.save(partial)
            print(f"부분 수집 {len(rows)}건 저장: {partial}")
        raise RuntimeError(f"크롤링 중 오류 발생 ({len(rows)}건까지 수집): {exc}") from exc
    return rows


def main():
    parser = argparse.ArgumentParser(description="MVNO 일일 요금제 및 가이드 위반 추적")
    parser.add_argument("--template", type=Path, default=TEMPLATE_FILE)
    parser.add_argument("--previous", type=Path, help="직전 정상 수집 파일 (기본값: 전일 날짜 파일)")
    parser.add_argument("--output", type=Path, help="결과 엑셀 경로 (하루 여러 번 실행 시 시각별 파일명 지정)")
    parser.add_argument("--from-xlsx", type=Path, help="브라우저 실행 없이 기존 파일로 보고서 작성")
    parser.add_argument("--date", type=date.fromisoformat, default=date.today())
    args = parser.parse_args()
    today = args.date.strftime("%Y%m%d")
    output = args.output or Path(__file__).with_name(f"mvno_combined_plans_{today}.xlsx")
    output.parent.mkdir(parents=True, exist_ok=True)
    previous = args.previous or output.with_name(f"mvno_combined_plans_{(args.date - timedelta(days=1)):%Y%m%d}.xlsx")
    if not previous.exists():
        if args.previous:
            raise FileNotFoundError(previous)
        previous = None
    if not args.template.exists():
        raise FileNotFoundError(f"템플릿 파일을 스크립트와 같은 폴더에 두세요: {args.template}")
    if args.from_xlsx:
        rows = sheet_rows(args.from_xlsx)
    else:
        # 실패한 수집 결과로 날짜별 정상 스냅샷을 덮어쓰지 않는다.
        rows = crawl_all()
    write_report(rows, args.template, previous, output)


if __name__ == "__main__":
    main()
